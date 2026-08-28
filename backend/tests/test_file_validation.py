import io
import zipfile

import docx
import openpyxl

from backend.app.services import file_validation
from backend.app.services.file_validation import (
    is_valid_docx,
    is_valid_pdf,
    is_valid_xlsx,
    validate_file_content,
    validate_safe_office_zip,
)


def _real_pdf_bytes() -> bytes:
    return b"%PDF-1.4\n%mock pdf content\n%%EOF"


def _real_docx_bytes() -> bytes:
    document = docx.Document()
    document.add_paragraph("hello")

    buffer = io.BytesIO()
    document.save(buffer)

    return buffer.getvalue()


def _real_xlsx_bytes() -> bytes:
    workbook = openpyxl.Workbook()
    workbook.active["A1"] = "hello"

    buffer = io.BytesIO()
    workbook.save(buffer)

    return buffer.getvalue()


def test_is_valid_pdf_accepts_real_pdf_signature():
    assert is_valid_pdf(_real_pdf_bytes()) is True


def test_is_valid_pdf_rejects_plain_text():
    assert is_valid_pdf(b"just some plain text") is False


def test_is_valid_docx_accepts_real_docx():
    assert is_valid_docx(_real_docx_bytes()) is True


def test_is_valid_xlsx_accepts_real_xlsx():
    assert is_valid_xlsx(_real_xlsx_bytes()) is True


def test_is_valid_docx_rejects_random_zip():
    buffer = io.BytesIO()

    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("hello.txt", "not a real docx")

    assert is_valid_docx(buffer.getvalue()) is False


def test_is_valid_xlsx_rejects_random_zip():
    buffer = io.BytesIO()

    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("hello.txt", "not a real xlsx")

    assert is_valid_xlsx(buffer.getvalue()) is False


def test_is_valid_docx_rejects_xlsx_bytes():
    assert is_valid_docx(_real_xlsx_bytes()) is False


def test_is_valid_xlsx_rejects_docx_bytes():
    assert is_valid_xlsx(_real_docx_bytes()) is False


def test_is_valid_docx_rejects_malformed_zip_without_raising():
    assert is_valid_docx(b"PK\x03\x04not a real zip stream") is False


def test_is_valid_xlsx_rejects_malformed_zip_without_raising():
    assert is_valid_xlsx(b"PK\x03\x04not a real zip stream") is False


def test_is_valid_docx_rejects_empty_bytes():
    assert is_valid_docx(b"") is False


def test_validate_file_content_dispatches_by_extension():
    assert (
        validate_file_content(".pdf", _real_pdf_bytes()) is True
    )
    assert (
        validate_file_content(".docx", _real_docx_bytes()) is True
    )
    assert (
        validate_file_content(".xlsx", _real_xlsx_bytes()) is True
    )


def test_validate_file_content_rejects_pdf_bytes_as_docx():
    assert (
        validate_file_content(".docx", _real_pdf_bytes()) is False
    )


def test_validate_file_content_rejects_unknown_extension():
    assert (
        validate_file_content(".txt", _real_pdf_bytes()) is False
    )


# --- ZIP bomb / decompression-abuse guard ---------------------------


def _make_zip(entries):
    buffer = io.BytesIO()

    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in entries:
            archive.writestr(name, data)

    return buffer.getvalue()


def test_validate_safe_office_zip_accepts_real_docx():
    assert validate_safe_office_zip(_real_docx_bytes()) is True


def test_validate_safe_office_zip_accepts_real_xlsx():
    assert validate_safe_office_zip(_real_xlsx_bytes()) is True


def test_validate_safe_office_zip_accepts_empty_entry():
    content = _make_zip([("empty.txt", b"")])

    assert validate_safe_office_zip(content) is True


def test_validate_safe_office_zip_rejects_malformed_zip_without_raising():
    assert (
        validate_safe_office_zip(b"PK\x03\x04not a real zip stream")
        is False
    )


def test_validate_safe_office_zip_rejects_excessive_entry_count(
    monkeypatch,
):
    monkeypatch.setattr(file_validation, "MAX_ZIP_ENTRIES", 2)

    content = _make_zip(
        [
            ("a.txt", b"x"),
            ("b.txt", b"y"),
            ("c.txt", b"z"),
        ]
    )

    assert validate_safe_office_zip(content) is False


def test_validate_safe_office_zip_rejects_single_entry_size(
    monkeypatch,
):
    monkeypatch.setattr(
        file_validation, "MAX_ENTRY_UNCOMPRESSED_SIZE", 50
    )

    content = _make_zip([("big.txt", b"D" * 100)])

    assert validate_safe_office_zip(content) is False


def test_validate_safe_office_zip_rejects_total_uncompressed_size(
    monkeypatch,
):
    monkeypatch.setattr(
        file_validation, "MAX_TOTAL_UNCOMPRESSED_SIZE", 100
    )

    content = _make_zip(
        [
            ("a.txt", b"B" * 60),
            ("b.txt", b"C" * 60),
        ]
    )

    assert validate_safe_office_zip(content) is False


def test_validate_safe_office_zip_rejects_excessive_compression_ratio(
    monkeypatch,
):
    monkeypatch.setattr(file_validation, "MAX_COMPRESSION_RATIO", 2)

    content = _make_zip([("big.txt", b"A" * 500)])

    assert validate_safe_office_zip(content) is False


def test_validate_safe_office_zip_rejects_compress_size_zero_edge_case(
    monkeypatch,
):
    fake_info = zipfile.ZipInfo("word/document.xml")
    fake_info.file_size = 1000
    fake_info.compress_size = 0

    monkeypatch.setattr(
        zipfile.ZipFile, "infolist", lambda self: [fake_info]
    )

    assert validate_safe_office_zip(_real_docx_bytes()) is False


def test_validate_safe_office_zip_rejects_parent_directory_traversal():
    content = _make_zip([("../evil.txt", b"data")])

    assert validate_safe_office_zip(content) is False


def test_validate_safe_office_zip_rejects_absolute_path_entry():
    content = _make_zip([("/etc/passwd", b"data")])

    assert validate_safe_office_zip(content) is False


def test_validate_safe_office_zip_rejects_windows_style_traversal():
    content = _make_zip([("..\\evil.txt", b"data")])

    assert validate_safe_office_zip(content) is False


def test_validate_safe_office_zip_rejects_windows_drive_absolute_path():
    content = _make_zip([("C:\\evil.txt", b"data")])

    assert validate_safe_office_zip(content) is False


def test_validate_safe_office_zip_rejects_encrypted_entry(monkeypatch):
    fake_info = zipfile.ZipInfo("word/document.xml")
    fake_info.file_size = 10
    fake_info.compress_size = 10
    fake_info.flag_bits = 0x1

    monkeypatch.setattr(
        zipfile.ZipFile, "infolist", lambda self: [fake_info]
    )

    assert validate_safe_office_zip(_real_docx_bytes()) is False
