import io
import zipfile

import docx
import openpyxl
import pymupdf
from fastapi.testclient import TestClient

from backend.app.main import app
from backend.app.routes.documents import UPLOAD_ROOT
from backend.app.services import file_validation

PDF_CONTENT_TYPE = "application/pdf"
DOCX_CONTENT_TYPE = (
    "application/vnd.openxmlformats"
    "-officedocument.wordprocessingml.document"
)
XLSX_CONTENT_TYPE = (
    "application/vnd.openxmlformats"
    "-officedocument.spreadsheetml.sheet"
)


def _real_pdf_bytes() -> bytes:
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "hello")

    buffer = io.BytesIO()
    document.save(buffer)
    document.close()

    return buffer.getvalue()


def _real_docx_bytes() -> bytes:
    document = docx.Document()
    document.add_paragraph("hello")

    buffer = io.BytesIO()
    document.save(buffer)

    return buffer.getvalue()


def _real_xlsx_bytes() -> bytes:
    workbook = openpyxl.Workbook()
    workbook.active["A1"] = "hello"

    buffer = io.BytesIO()
    workbook.save(buffer)

    return buffer.getvalue()


def _random_zip_bytes() -> bytes:
    buffer = io.BytesIO()

    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("readme.txt", "just a random zip")

    return buffer.getvalue()


def _docx_like_zip(extra_entries=()):
    buffer = io.BytesIO()

    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("word/document.xml", "<document/>")

        for name, data in extra_entries:
            archive.writestr(name, data)

    return buffer.getvalue()


def _upload(client, filename, content, content_type):
    return client.post(
        "/api/documents",
        files={
            "file": (
                filename,
                io.BytesIO(content),
                content_type,
            )
        },
    )


def _existing_document_ids():
    if not UPLOAD_ROOT.exists():
        return set()

    return {
        entry.name
        for entry in UPLOAD_ROOT.iterdir()
        if entry.is_dir()
    }


def test_upload_accepts_valid_pdf():
    client = TestClient(app)

    response = _upload(
        client,
        "sample.pdf",
        _real_pdf_bytes(),
        PDF_CONTENT_TYPE,
    )

    assert response.status_code == 200
    assert response.json()["status"] == "uploaded"


def test_upload_rejects_fake_pdf_plain_text():
    client = TestClient(app)

    before = _existing_document_ids()

    response = _upload(
        client,
        "fake.pdf",
        b"this is just plain text, not a pdf",
        PDF_CONTENT_TYPE,
    )

    assert response.status_code == 415
    assert _existing_document_ids() == before


def test_upload_rejects_pdf_bytes_renamed_to_docx():
    client = TestClient(app)

    before = _existing_document_ids()

    response = _upload(
        client,
        "fake.docx",
        _real_pdf_bytes(),
        DOCX_CONTENT_TYPE,
    )

    assert response.status_code == 415
    assert _existing_document_ids() == before


def test_upload_accepts_valid_docx():
    client = TestClient(app)

    response = _upload(
        client,
        "sample.docx",
        _real_docx_bytes(),
        DOCX_CONTENT_TYPE,
    )

    assert response.status_code == 200
    assert response.json()["status"] == "uploaded"


def test_upload_rejects_random_zip_renamed_to_docx():
    client = TestClient(app)

    before = _existing_document_ids()

    response = _upload(
        client,
        "fake.docx",
        _random_zip_bytes(),
        DOCX_CONTENT_TYPE,
    )

    assert response.status_code == 415
    assert _existing_document_ids() == before


def test_upload_accepts_valid_xlsx():
    client = TestClient(app)

    response = _upload(
        client,
        "sample.xlsx",
        _real_xlsx_bytes(),
        XLSX_CONTENT_TYPE,
    )

    assert response.status_code == 200
    assert response.json()["status"] == "uploaded"


def test_upload_rejects_docx_renamed_to_xlsx():
    client = TestClient(app)

    before = _existing_document_ids()

    response = _upload(
        client,
        "fake.xlsx",
        _real_docx_bytes(),
        XLSX_CONTENT_TYPE,
    )

    assert response.status_code == 415
    assert _existing_document_ids() == before


def test_upload_rejects_xlsx_renamed_to_docx():
    client = TestClient(app)

    before = _existing_document_ids()

    response = _upload(
        client,
        "fake.docx",
        _real_xlsx_bytes(),
        DOCX_CONTENT_TYPE,
    )

    assert response.status_code == 415
    assert _existing_document_ids() == before


def test_upload_rejects_malformed_zip_cleanly():
    client = TestClient(app)

    before = _existing_document_ids()

    response = _upload(
        client,
        "fake.docx",
        b"PK\x03\x04" + b"\x00" * 32,
        DOCX_CONTENT_TYPE,
    )

    assert response.status_code == 415
    assert "detail" in response.json()
    assert _existing_document_ids() == before


def test_upload_rejects_when_declared_mime_wrong_even_if_bytes_valid():
    """Declared Content-Type is checked first (existing early/UX
    rejection); real, otherwise-valid PDF bytes sent with the wrong
    declared MIME are still rejected per current policy.
    """
    client = TestClient(app)

    response = _upload(
        client,
        "sample.pdf",
        _real_pdf_bytes(),
        "text/plain",
    )

    assert response.status_code == 415


def test_rejected_upload_leaves_no_source_document_on_disk():
    client = TestClient(app)

    before = _existing_document_ids()

    response = _upload(
        client,
        "fake.pdf",
        b"not a pdf at all",
        PDF_CONTENT_TYPE,
    )

    assert response.status_code == 415
    assert _existing_document_ids() == before


def test_upload_rejects_docx_exceeding_zip_entry_count_guard(
    monkeypatch,
):
    monkeypatch.setattr(file_validation, "MAX_ZIP_ENTRIES", 2)

    client = TestClient(app)
    before = _existing_document_ids()

    content = _docx_like_zip(
        extra_entries=[("extra1.txt", "x"), ("extra2.txt", "y")]
    )

    response = _upload(
        client,
        "sample.docx",
        content,
        DOCX_CONTENT_TYPE,
    )

    assert response.status_code == 413
    assert _existing_document_ids() == before


def test_upload_rejects_docx_exceeding_total_uncompressed_size_guard(
    monkeypatch,
):
    monkeypatch.setattr(
        file_validation, "MAX_TOTAL_UNCOMPRESSED_SIZE", 100
    )

    client = TestClient(app)
    before = _existing_document_ids()

    content = _docx_like_zip(
        extra_entries=[("big.txt", "A" * 200)]
    )

    response = _upload(
        client,
        "sample.docx",
        content,
        DOCX_CONTENT_TYPE,
    )

    assert response.status_code == 413
    assert _existing_document_ids() == before


def test_upload_rejects_docx_exceeding_single_entry_size_guard(
    monkeypatch,
):
    monkeypatch.setattr(
        file_validation, "MAX_ENTRY_UNCOMPRESSED_SIZE", 50
    )

    client = TestClient(app)
    before = _existing_document_ids()

    content = _docx_like_zip(
        extra_entries=[("big.txt", "A" * 100)]
    )

    response = _upload(
        client,
        "sample.docx",
        content,
        DOCX_CONTENT_TYPE,
    )

    assert response.status_code == 413
    assert _existing_document_ids() == before


def test_upload_rejects_docx_exceeding_compression_ratio_guard(
    monkeypatch,
):
    monkeypatch.setattr(file_validation, "MAX_COMPRESSION_RATIO", 2)

    client = TestClient(app)
    before = _existing_document_ids()

    content = _docx_like_zip(
        extra_entries=[("big.txt", "A" * 2000)]
    )

    response = _upload(
        client,
        "sample.docx",
        content,
        DOCX_CONTENT_TYPE,
    )

    assert response.status_code == 413
    assert _existing_document_ids() == before


def test_upload_rejects_docx_with_path_traversal_entry():
    client = TestClient(app)
    before = _existing_document_ids()

    content = _docx_like_zip(
        extra_entries=[("../evil.txt", "data")]
    )

    response = _upload(
        client,
        "sample.docx",
        content,
        DOCX_CONTENT_TYPE,
    )

    assert response.status_code == 413
    assert _existing_document_ids() == before


def test_upload_still_enforces_max_file_size():
    client = TestClient(app)

    oversized = _real_pdf_bytes() + (b"0" * (20 * 1024 * 1024))

    before = _existing_document_ids()

    response = _upload(
        client,
        "big.pdf",
        oversized,
        PDF_CONTENT_TYPE,
    )

    assert response.status_code == 413
    assert _existing_document_ids() == before
